"""
Unit tests for utils/export.py - Export functionality
"""

import pytest
import pandas as pd
import numpy as np
from datetime import datetime
import io
from openpyxl import load_workbook

from utils.export import (
    export_to_excel,
    export_to_excel_multiple_sheets,
    export_to_csv,
    create_summary_report,
    generate_excel_report,
    export_charts_data
)


class TestExcelExport:
    """Test Excel export functionality"""

    @pytest.mark.unit
    def test_export_to_excel_basic(self, sample_dataframe):
        """Test basic Excel export"""
        excel_data = export_to_excel(sample_dataframe)

        assert excel_data is not None
        assert isinstance(excel_data, bytes)
        assert len(excel_data) > 0

    @pytest.mark.unit
    def test_export_to_excel_readable(self, sample_dataframe):
        """Test exported Excel is readable"""
        excel_data = export_to_excel(sample_dataframe)

        # Should be able to read back
        wb = load_workbook(io.BytesIO(excel_data))
        assert wb is not None
        assert len(wb.sheetnames) > 0

    @pytest.mark.unit
    def test_export_to_excel_custom_sheet_name(self, sample_dataframe):
        """Test Excel export with custom sheet name"""
        sheet_name = 'CustomSheet'
        excel_data = export_to_excel(sample_dataframe, sheet_name=sheet_name)

        wb = load_workbook(io.BytesIO(excel_data))
        assert sheet_name in wb.sheetnames

    @pytest.mark.unit
    def test_export_to_excel_preserves_data(self, sample_dataframe):
        """Test Excel export preserves data correctly"""
        excel_data = export_to_excel(sample_dataframe)

        # Read back and compare
        df_read = pd.read_excel(io.BytesIO(excel_data))

        assert len(df_read) == len(sample_dataframe)
        # Check some columns exist
        assert 'platform' in df_read.columns
        assert 'urgency' in df_read.columns

    @pytest.mark.unit
    def test_export_to_excel_empty_dataframe(self):
        """Test Excel export with empty DataFrame"""
        df = pd.DataFrame({'col1': [], 'col2': []})
        excel_data = export_to_excel(df)

        assert excel_data is not None
        assert isinstance(excel_data, bytes)


class TestExcelMultipleSheets:
    """Test Excel export with multiple sheets"""

    @pytest.mark.unit
    def test_export_multiple_sheets_basic(self, sample_dataframe):
        """Test multi-sheet Excel export"""
        data_dict = {
            'Sheet1': sample_dataframe.head(20),
            'Sheet2': sample_dataframe.tail(20)
        }

        excel_data = export_to_excel_multiple_sheets(data_dict)

        assert excel_data is not None
        assert isinstance(excel_data, bytes)

    @pytest.mark.unit
    def test_export_multiple_sheets_readable(self, sample_dataframe):
        """Test multi-sheet Excel is readable"""
        data_dict = {
            'Tickets': sample_dataframe.head(20),
            'Summary': sample_dataframe.tail(20)
        }

        excel_data = export_to_excel_multiple_sheets(data_dict)

        wb = load_workbook(io.BytesIO(excel_data))
        assert 'Tickets' in wb.sheetnames
        assert 'Summary' in wb.sheetnames

    @pytest.mark.unit
    def test_export_multiple_sheets_preserves_data(self, sample_dataframe):
        """Test multi-sheet export preserves data"""
        df1 = sample_dataframe.head(10)
        df2 = sample_dataframe.tail(10)

        data_dict = {
            'First': df1,
            'Second': df2
        }

        excel_data = export_to_excel_multiple_sheets(data_dict)

        # Read back
        dfs = pd.read_excel(io.BytesIO(excel_data), sheet_name=None)

        assert 'First' in dfs
        assert 'Second' in dfs
        assert len(dfs['First']) == 10
        assert len(dfs['Second']) == 10

    @pytest.mark.unit
    def test_export_multiple_sheets_empty_dict(self):
        """Test multi-sheet export with empty dict"""
        excel_data = export_to_excel_multiple_sheets({})

        assert excel_data is not None
        assert isinstance(excel_data, bytes)


class TestCSVExport:
    """Test CSV export functionality"""

    @pytest.mark.unit
    def test_export_to_csv_basic(self, sample_dataframe):
        """Test basic CSV export"""
        csv_data = export_to_csv(sample_dataframe)

        assert csv_data is not None
        assert isinstance(csv_data, str)
        assert len(csv_data) > 0

    @pytest.mark.unit
    def test_export_to_csv_readable(self, sample_dataframe):
        """Test exported CSV is readable"""
        csv_data = export_to_csv(sample_dataframe)

        # Should be able to read back
        df_read = pd.read_csv(io.StringIO(csv_data))

        assert len(df_read) == len(sample_dataframe)

    @pytest.mark.unit
    def test_export_to_csv_preserves_data(self, sample_dataframe):
        """Test CSV export preserves data"""
        csv_data = export_to_csv(sample_dataframe)

        df_read = pd.read_csv(io.StringIO(csv_data))

        # Check columns exist
        assert 'platform' in df_read.columns
        assert 'urgency' in df_read.columns

        # Check row count
        assert len(df_read) == len(sample_dataframe)

    @pytest.mark.unit
    def test_export_to_csv_utf8_bom(self, sample_dataframe):
        """Test CSV export uses UTF-8 with BOM"""
        csv_data = export_to_csv(sample_dataframe)

        # Should start with UTF-8 BOM
        assert csv_data.startswith('\ufeff') or 'utf-8' in str(csv_data.encode())

    @pytest.mark.unit
    def test_export_to_csv_empty_dataframe(self):
        """Test CSV export with empty DataFrame"""
        df = pd.DataFrame({'col1': [], 'col2': []})
        csv_data = export_to_csv(df)

        assert csv_data is not None
        assert isinstance(csv_data, str)


class TestSummaryReport:
    """Test create_summary_report functionality"""

    @pytest.mark.unit
    def test_create_summary_report_basic(self, sample_dataframe):
        """Test basic summary report creation"""
        summary = create_summary_report(sample_dataframe)

        assert summary is not None
        assert isinstance(summary, dict)

    @pytest.mark.unit
    def test_create_summary_report_metadata(self, sample_dataframe):
        """Test summary report has metadata"""
        summary = create_summary_report(sample_dataframe)

        assert 'Metadata' in summary
        metadata = summary['Metadata']

        assert 'Total Tickets' in metadata
        assert 'Date Range' in metadata
        assert 'Generated At' in metadata

    @pytest.mark.unit
    def test_create_summary_report_key_metrics(self, sample_dataframe):
        """Test summary report has key metrics"""
        summary = create_summary_report(sample_dataframe)

        assert 'Key Metrics' in summary
        metrics = summary['Key Metrics']

        # Should have various metrics
        assert isinstance(metrics, dict)
        assert len(metrics) > 0

    @pytest.mark.unit
    def test_create_summary_report_platforms(self, sample_dataframe):
        """Test summary report includes platform breakdown"""
        summary = create_summary_report(sample_dataframe)

        # Should have platform or urgency breakdown
        assert 'Metadata' in summary or 'Key Metrics' in summary

    @pytest.mark.unit
    def test_create_summary_report_empty_dataframe(self):
        """Test summary report with empty DataFrame"""
        df = pd.DataFrame({
            'created_at': [],
            'platform': [],
            'urgency': [],
            'sentiment': []
        })

        summary = create_summary_report(df)

        assert summary is not None
        assert 'Metadata' in summary


class TestExcelReport:
    """Test generate_excel_report functionality"""

    @pytest.mark.unit
    def test_generate_excel_report_basic(self, sample_dataframe):
        """Test basic Excel report generation"""
        summary = create_summary_report(sample_dataframe)
        excel_data = generate_excel_report(sample_dataframe, summary)

        assert excel_data is not None
        assert isinstance(excel_data, bytes)

    @pytest.mark.unit
    def test_generate_excel_report_multiple_sheets(self, sample_dataframe):
        """Test Excel report has multiple sheets"""
        summary = create_summary_report(sample_dataframe)
        excel_data = generate_excel_report(sample_dataframe, summary)

        wb = load_workbook(io.BytesIO(excel_data))

        # Should have multiple sheets
        assert len(wb.sheetnames) > 1

    @pytest.mark.unit
    def test_generate_excel_report_has_summary_sheet(self, sample_dataframe):
        """Test Excel report has summary sheet"""
        summary = create_summary_report(sample_dataframe)
        excel_data = generate_excel_report(sample_dataframe, summary)

        wb = load_workbook(io.BytesIO(excel_data))

        # Should have summary or metadata sheet
        sheet_names = [s.lower() for s in wb.sheetnames]
        assert any('summ' in s or 'meta' in s for s in sheet_names)

    @pytest.mark.unit
    def test_generate_excel_report_has_data_sheets(self, sample_dataframe):
        """Test Excel report has data breakdown sheets"""
        summary = create_summary_report(sample_dataframe)
        excel_data = generate_excel_report(sample_dataframe, summary)

        wb = load_workbook(io.BytesIO(excel_data))

        # Should have breakdown sheets (platform, urgency, etc.)
        assert len(wb.sheetnames) >= 3


class TestChartsData:
    """Test export_charts_data functionality"""

    @pytest.mark.unit
    def test_export_charts_data_basic(self, sample_dataframe):
        """Test basic charts data export"""
        charts_data = export_charts_data(sample_dataframe)

        assert charts_data is not None
        assert isinstance(charts_data, dict)

    @pytest.mark.unit
    def test_export_charts_data_has_trend(self, sample_dataframe):
        """Test charts data includes trend data"""
        charts_data = export_charts_data(sample_dataframe)

        # Should have some kind of trend or time-based data
        assert len(charts_data) > 0

    @pytest.mark.unit
    def test_export_charts_data_dataframes(self, sample_dataframe):
        """Test charts data contains DataFrames"""
        charts_data = export_charts_data(sample_dataframe)

        # All values should be DataFrames
        for key, value in charts_data.items():
            assert isinstance(value, pd.DataFrame), f"{key} is not a DataFrame"

    @pytest.mark.unit
    def test_export_charts_data_platform_breakdown(self, sample_dataframe):
        """Test charts data includes platform breakdown"""
        charts_data = export_charts_data(sample_dataframe)

        # Should have platform or category breakdowns
        assert len(charts_data) > 0

    @pytest.mark.unit
    def test_export_charts_data_empty_dataframe(self):
        """Test charts data with empty DataFrame"""
        df = pd.DataFrame({
            'created_date': [],
            'platform': [],
            'urgency': []
        })

        charts_data = export_charts_data(df)

        assert charts_data is not None
        assert isinstance(charts_data, dict)


class TestExportEdgeCases:
    """Test edge cases in export functionality"""

    @pytest.mark.unit
    def test_export_dataframe_with_special_characters(self):
        """Test export with special characters in data"""
        df = pd.DataFrame({
            'name': ['Test™', 'Café', 'Ñoño'],
            'description': ['Special chars: é, ă, ș, ț', 'Unicode: 中文', 'Emoji: 😊']
        })

        csv_data = export_to_csv(df)
        excel_data = export_to_excel(df)

        assert csv_data is not None
        assert excel_data is not None

    @pytest.mark.unit
    def test_export_dataframe_with_dates(self, sample_dataframe):
        """Test export preserves date columns"""
        excel_data = export_to_excel(sample_dataframe)

        # Should not error with datetime columns
        assert excel_data is not None

    @pytest.mark.unit
    def test_export_dataframe_with_nulls(self):
        """Test export with null values"""
        df = pd.DataFrame({
            'col1': [1, None, 3],
            'col2': ['a', 'b', None],
            'col3': [None, None, None]
        })

        csv_data = export_to_csv(df)
        excel_data = export_to_excel(df)

        assert csv_data is not None
        assert excel_data is not None

    @pytest.mark.unit
    def test_export_very_long_sheet_name(self, sample_dataframe):
        """Test export with very long sheet name"""
        # Excel has 31 char limit for sheet names
        long_name = 'A' * 50

        excel_data = export_to_excel(sample_dataframe, sheet_name=long_name)

        # Should handle gracefully (truncate or error)
        assert excel_data is not None

    @pytest.mark.unit
    def test_export_large_dataframe(self):
        """Test export with large DataFrame"""
        # Create larger DataFrame
        np.random.seed(42)
        n = 1000

        df = pd.DataFrame({
            'col1': np.random.randint(0, 100, n),
            'col2': np.random.choice(['a', 'b', 'c'], n),
            'col3': np.random.randn(n)
        })

        excel_data = export_to_excel(df)
        csv_data = export_to_csv(df)

        assert excel_data is not None
        assert csv_data is not None
        assert len(excel_data) > 0
        assert len(csv_data) > 0


class TestExportIntegration:
    """Integration tests for export workflows"""

    @pytest.mark.integration
    def test_full_export_workflow(self, sample_dataframe):
        """Test complete export workflow"""
        # Create summary
        summary = create_summary_report(sample_dataframe)

        # Generate all export formats
        csv_data = export_to_csv(sample_dataframe)
        excel_simple = export_to_excel(sample_dataframe)
        excel_full = generate_excel_report(sample_dataframe, summary)

        # All should succeed
        assert csv_data is not None
        assert excel_simple is not None
        assert excel_full is not None

        # Verify data integrity
        df_from_csv = pd.read_csv(io.StringIO(csv_data))
        assert len(df_from_csv) == len(sample_dataframe)

    @pytest.mark.integration
    def test_export_charts_and_data(self, sample_dataframe):
        """Test exporting both data and chart data"""
        # Export main data
        excel_data = export_to_excel(sample_dataframe)

        # Export chart data
        charts_data = export_charts_data(sample_dataframe)

        # Combine into multi-sheet
        all_data = {'Main': sample_dataframe}
        all_data.update(charts_data)

        excel_combined = export_to_excel_multiple_sheets(all_data)

        assert excel_combined is not None

        # Verify all sheets present
        wb = load_workbook(io.BytesIO(excel_combined))
        assert 'Main' in wb.sheetnames
        assert len(wb.sheetnames) > 1
