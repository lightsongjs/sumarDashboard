"""
Integration tests for SmartBill Support Analytics
Tests full workflows from CSV loading to export
"""

import pytest
import pandas as pd
import numpy as np
from datetime import datetime, timedelta
import io

from utils.data_processor import CSVProcessor
from utils.charts import (
    create_trend_chart,
    create_platform_distribution,
    create_urgency_trend
)
from utils.export import (
    export_to_excel,
    export_to_csv,
    create_summary_report,
    generate_excel_report
)


class TestEndToEndWorkflow:
    """Test complete end-to-end workflows"""

    @pytest.mark.integration
    def test_csv_to_excel_workflow(self, sample_clean_csv):
        """Test complete workflow: Load CSV → Process → Export Excel"""
        # 1. Load CSV
        df = CSVProcessor.load_csv(sample_clean_csv)
        assert df is not None

        # 2. Process
        df_processed = CSVProcessor.process_dataframe(df)
        assert df_processed is not None
        assert len(df_processed) > 0

        # 3. Validate
        is_valid, errors = CSVProcessor.validate_dataframe(df_processed)
        assert is_valid is True

        # 4. Export
        excel_data = export_to_excel(df_processed)
        assert excel_data is not None
        assert len(excel_data) > 0

    @pytest.mark.integration
    def test_csv_filter_export_workflow(self, sample_clean_csv):
        """Test workflow: Load → Filter → Export"""
        # Load and process
        df = CSVProcessor.load_csv(sample_clean_csv)
        df_processed = CSVProcessor.process_dataframe(df)

        # Apply filters
        end_date = datetime.now()
        start_date = end_date - timedelta(days=30)

        df_filtered = CSVProcessor.apply_filters(
            df_processed,
            date_range=(start_date, end_date),
            platforms=['facturare'],
            urgency_levels=['blocker']
        )

        # Export filtered data
        csv_data = export_to_csv(df_filtered)
        assert csv_data is not None

        # Verify filtered data can be read back
        df_from_csv = pd.read_csv(io.StringIO(csv_data))
        assert len(df_from_csv) <= len(df_processed)

    @pytest.mark.integration
    def test_csv_charts_workflow(self, sample_clean_csv):
        """Test workflow: Load → Process → Generate Charts"""
        # Load and process
        df = CSVProcessor.load_csv(sample_clean_csv)
        df_processed = CSVProcessor.process_dataframe(df)

        # Generate various charts
        fig_trend = create_trend_chart(df_processed)
        fig_platform = create_platform_distribution(df_processed)
        fig_urgency = create_urgency_trend(df_processed)

        # All charts should be generated successfully
        assert fig_trend is not None
        assert fig_platform is not None
        assert fig_urgency is not None

    @pytest.mark.integration
    def test_full_report_generation_workflow(self, sample_clean_csv):
        """Test complete report generation workflow"""
        # 1. Load and process
        df = CSVProcessor.load_csv(sample_clean_csv)
        df_processed = CSVProcessor.process_dataframe(df)

        # 2. Create summary
        summary = create_summary_report(df_processed)
        assert summary is not None
        assert 'Metadata' in summary

        # 3. Generate full Excel report
        excel_report = generate_excel_report(df_processed, summary)
        assert excel_report is not None

        # 4. Verify report is readable
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(excel_report))
        assert len(wb.sheetnames) > 1


class TestDirtyDataWorkflow:
    """Test workflows with dirty/messy CSV data"""

    @pytest.mark.integration
    @pytest.mark.csv
    def test_dirty_csv_complete_workflow(self, sample_dirty_csv):
        """Test complete workflow with dirty CSV data"""
        # Load dirty CSV
        df = CSVProcessor.load_csv(sample_dirty_csv)

        # Should remove empty first column
        assert ' ' not in df.columns

        # Process (normalize lowercase, etc.)
        df_processed = CSVProcessor.process_dataframe(df)

        # Verify normalization worked
        assert df_processed['platform'].iloc[0] == 'facturare'  # Was 'Facturare'

        # Apply filters
        df_filtered = CSVProcessor.apply_filters(
            df_processed,
            urgency_levels=['blocker']
        )

        # Should find blockers despite mixed case in original
        assert len(df_filtered) > 0

        # Export should work
        excel_data = export_to_excel(df_filtered)
        assert excel_data is not None

    @pytest.mark.integration
    def test_dirty_csv_resilience(self, sample_dirty_csv):
        """Test system resilience with messy data"""
        # Load
        df = CSVProcessor.load_csv(sample_dirty_csv)

        # Process (should handle mixed case, missing values, etc.)
        df_processed = CSVProcessor.process_dataframe(df)

        # Should have computed columns despite messy input
        assert 'priority_score' in df_processed.columns
        assert 'created_date' in df_processed.columns

        # Charts should still generate
        fig = create_trend_chart(df_processed)
        assert fig is not None

        # Export should work
        csv_data = export_to_csv(df_processed)
        assert csv_data is not None


class TestMultipleFiltersWorkflow:
    """Test workflows with multiple filters combined"""

    @pytest.mark.integration
    def test_progressive_filtering(self, sample_clean_csv):
        """Test applying filters progressively"""
        df = CSVProcessor.load_csv(sample_clean_csv)
        df_processed = CSVProcessor.process_dataframe(df)

        original_count = len(df_processed)

        # Filter 1: Date range
        end_date = datetime.now()
        start_date = end_date - timedelta(days=30)
        df_filtered = CSVProcessor.apply_filters(
            df_processed,
            date_range=(start_date, end_date)
        )
        count_after_date = len(df_filtered)
        assert count_after_date <= original_count

        # Filter 2: Add platform filter
        df_filtered = CSVProcessor.apply_filters(
            df_filtered,
            platforms=['facturare', 'gestiune']
        )
        count_after_platform = len(df_filtered)
        assert count_after_platform <= count_after_date

        # Filter 3: Add urgency filter
        df_filtered = CSVProcessor.apply_filters(
            df_filtered,
            urgency_levels=['blocker']
        )
        count_after_urgency = len(df_filtered)
        assert count_after_urgency <= count_after_platform

    @pytest.mark.integration
    def test_complex_filter_combination(self, sample_dataframe):
        """Test complex combination of all filters"""
        end_date = datetime.now()
        start_date = end_date - timedelta(days=60)

        df_filtered = CSVProcessor.apply_filters(
            sample_dataframe,
            date_range=(start_date, end_date),
            platforms=['facturare'],
            urgency_levels=['blocker', 'mediu'],
            sentiment='negativ',
            only_recurrent=True,
            affects_business=True
        )

        # Should apply all filters correctly
        if len(df_filtered) > 0:
            assert all(df_filtered['platform'] == 'facturare')
            assert all(df_filtered['sentiment'] == 'negativ')
            assert all(df_filtered['is_recurrent'] == True)
            assert all(df_filtered['affects_business_flow'] == True)


class TestDataIntegrity:
    """Test data integrity through processing pipeline"""

    @pytest.mark.integration
    def test_data_preservation_through_pipeline(self, sample_clean_csv):
        """Test data is preserved correctly through pipeline"""
        # Load
        df_original = CSVProcessor.load_csv(sample_clean_csv)
        original_rows = len(df_original)

        # Process
        df_processed = CSVProcessor.process_dataframe(df_original)

        # Should have same number of rows
        assert len(df_processed) == original_rows

        # Should preserve original data columns
        assert 'platform' in df_processed.columns
        assert 'urgency' in df_processed.columns

    @pytest.mark.integration
    def test_computed_columns_consistency(self, sample_clean_csv):
        """Test computed columns are consistent"""
        df = CSVProcessor.load_csv(sample_clean_csv)
        df_processed = CSVProcessor.process_dataframe(df)

        # Check priority_score consistency
        for idx, row in df_processed.iterrows():
            expected_score = 0

            if row['urgency'] == 'blocker':
                expected_score += 3
            elif row['urgency'] == 'mediu':
                expected_score += 2
            elif row['urgency'] == 'scazut':
                expected_score += 1

            if row['affects_business_flow']:
                expected_score += 2

            if row['is_recurrent']:
                expected_score += 1

            assert row['priority_score'] == expected_score

    @pytest.mark.integration
    def test_date_calculations_accuracy(self, sample_clean_csv):
        """Test date-based calculations are accurate"""
        df = CSVProcessor.load_csv(sample_clean_csv)
        df_processed = CSVProcessor.process_dataframe(df)

        # Check resolution time for closed tickets
        for idx, row in df_processed.iterrows():
            if pd.notna(row['closed_at']) and pd.notna(row['created_at']):
                # Calculate manually
                time_diff = (row['closed_at'] - row['created_at']).total_seconds() / 3600

                # Should match computed column
                assert row['resolution_time_hours'] == pytest.approx(time_diff, rel=0.01)


class TestExportIntegrity:
    """Test export maintains data integrity"""

    @pytest.mark.integration
    def test_csv_export_roundtrip(self, sample_dataframe):
        """Test CSV export → import roundtrip"""
        # Export
        csv_data = export_to_csv(sample_dataframe)

        # Import back
        df_reimported = pd.read_csv(io.StringIO(csv_data))

        # Should have same shape
        assert len(df_reimported) == len(sample_dataframe)
        assert len(df_reimported.columns) == len(sample_dataframe.columns)

    @pytest.mark.integration
    def test_excel_export_roundtrip(self, sample_dataframe):
        """Test Excel export → import roundtrip"""
        # Export
        excel_data = export_to_excel(sample_dataframe)

        # Import back
        df_reimported = pd.read_excel(io.BytesIO(excel_data))

        # Should have same number of rows
        assert len(df_reimported) == len(sample_dataframe)

    @pytest.mark.integration
    def test_multi_sheet_excel_integrity(self, sample_dataframe):
        """Test multi-sheet Excel maintains data integrity"""
        from utils.export import export_to_excel_multiple_sheets

        # Create multiple sheets
        sheets = {
            'All': sample_dataframe,
            'Blockers': sample_dataframe[sample_dataframe['urgency'] == 'blocker'],
            'Positive': sample_dataframe[sample_dataframe['sentiment'] == 'pozitiv']
        }

        # Export
        excel_data = export_to_excel_multiple_sheets(sheets)

        # Import and verify
        dfs = pd.read_excel(io.BytesIO(excel_data), sheet_name=None)

        assert 'All' in dfs
        assert len(dfs['All']) == len(sample_dataframe)


class TestPerformanceIntegration:
    """Test performance with realistic data volumes"""

    @pytest.mark.integration
    @pytest.mark.slow
    def test_large_csv_processing(self):
        """Test processing large CSV files"""
        # Create large DataFrame
        np.random.seed(42)
        n = 5000

        now = datetime.now()
        df_large = pd.DataFrame({
            'created_at': [now - timedelta(days=np.random.randint(0, 90)) for _ in range(n)],
            'closed_at': [now - timedelta(days=np.random.randint(0, 85)) if i % 3 != 0 else pd.NaT for i in range(n)],
            'platform': np.random.choice(['facturare', 'gestiune', 'api', 'spv'], n),
            'area': np.random.choice(['SPV', 'Documente', 'Clienti', 'Rapoarte'], n),
            'urgency': np.random.choice(['blocker', 'mediu', 'scazut'], n),
            'sentiment': np.random.choice(['pozitiv', 'neutru', 'negativ'], n),
            'is_recurrent': np.random.choice([True, False], n),
            'affects_business_flow': np.random.choice([True, False], n),
            'problem_type': np.random.choice(['bug', 'feature_gap', 'usability'], n),
        })

        # Process
        df_processed = CSVProcessor.process_dataframe(df_large)

        # Should complete successfully
        assert len(df_processed) == n

        # Apply filters
        df_filtered = CSVProcessor.apply_filters(
            df_processed,
            platforms=['facturare'],
            urgency_levels=['blocker']
        )

        # Should filter correctly
        assert len(df_filtered) <= n

        # Export should work
        excel_data = export_to_excel(df_filtered)
        assert len(excel_data) > 0

    @pytest.mark.integration
    @pytest.mark.slow
    def test_multiple_chart_generation(self, sample_dataframe):
        """Test generating multiple charts doesn't cause issues"""
        from utils.charts import (
            create_sentiment_trend,
            create_problem_type_chart,
            create_heatmap
        )

        # Generate many charts
        charts = [
            create_trend_chart(sample_dataframe),
            create_platform_distribution(sample_dataframe),
            create_urgency_trend(sample_dataframe),
            create_sentiment_trend(sample_dataframe),
            create_problem_type_chart(sample_dataframe),
            create_heatmap(sample_dataframe)
        ]

        # All should succeed
        assert all(chart is not None for chart in charts)


class TestErrorRecovery:
    """Test error handling and recovery in workflows"""

    @pytest.mark.integration
    def test_invalid_csv_graceful_failure(self, tmp_path):
        """Test graceful failure with invalid CSV"""
        # Create invalid CSV
        invalid_csv = tmp_path / "invalid.csv"
        invalid_csv.write_text("not,valid,csv\ndata")

        # Should raise appropriate error
        with pytest.raises(Exception):
            df = CSVProcessor.load_csv(str(invalid_csv))
            CSVProcessor.validate_dataframe(df)

    @pytest.mark.integration
    def test_missing_columns_handling(self):
        """Test handling of missing optional columns"""
        # DataFrame with only critical columns
        df_minimal = pd.DataFrame({
            'created_at': ['2025-10-01 10:00:00', '2025-10-02 11:00:00'],
            'platform': ['facturare', 'gestiune'],
            'area': ['SPV', 'Docs']
        })

        # Process should work
        df_processed = CSVProcessor.process_dataframe(df_minimal)

        # Should add computed columns
        assert 'created_date' in df_processed.columns

        # Export should work
        csv_data = export_to_csv(df_processed)
        assert csv_data is not None

    @pytest.mark.integration
    def test_all_null_column_handling(self):
        """Test handling columns with all null values"""
        df = pd.DataFrame({
            'created_at': ['2025-10-01'] * 5,
            'platform': ['facturare'] * 5,
            'area': ['SPV'] * 5,
            'urgency': [None] * 5,  # All null
            'sentiment': [None] * 5  # All null
        })

        # Should process without errors
        df_processed = CSVProcessor.process_dataframe(df)

        # Should still export
        excel_data = export_to_excel(df_processed)
        assert excel_data is not None


class TestRealWorldScenarios:
    """Test realistic usage scenarios"""

    @pytest.mark.integration
    def test_weekly_report_generation(self, sample_dataframe):
        """Test generating weekly report"""
        # Filter for last 7 days
        end_date = datetime.now()
        start_date = end_date - timedelta(days=7)

        df_week = CSVProcessor.apply_filters(
            sample_dataframe,
            date_range=(start_date, end_date)
        )

        # Create summary
        summary = create_summary_report(df_week)

        # Generate report
        excel_report = generate_excel_report(df_week, summary)

        assert excel_report is not None

    @pytest.mark.integration
    def test_blocker_analysis_workflow(self, sample_dataframe):
        """Test workflow for analyzing blocker issues"""
        # Filter only blockers
        df_blockers = CSVProcessor.apply_filters(
            sample_dataframe,
            urgency_levels=['blocker']
        )

        # Should have priority scores
        if len(df_blockers) > 0:
            assert all(df_blockers['priority_score'] >= 3)

            # Export for analysis
            csv_data = export_to_csv(df_blockers)
            assert csv_data is not None

    @pytest.mark.integration
    def test_platform_comparison_workflow(self, sample_dataframe):
        """Test comparing metrics across platforms"""
        platforms = sample_dataframe['platform'].unique()

        results = {}
        for platform in platforms:
            df_platform = CSVProcessor.apply_filters(
                sample_dataframe,
                platforms=[platform]
            )

            results[platform] = {
                'total': len(df_platform),
                'blockers': len(df_platform[df_platform['urgency'] == 'blocker'])
            }

        # Should have results for each platform
        assert len(results) > 0
        assert all('total' in v for v in results.values())
